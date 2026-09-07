"""Phase 4 W5 锁文件验收脚本：直接测 _safe_save 在文件被占用时的回退行为

验收：_safe_save 收到 PermissionError → 自动创建 _takeoff/ 目录 →
     文件保存到 <原目录>/_takeoff/<文件名>，返回该路径。
"""
import ctypes
import os
import shutil
import sys
import tempfile

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)) + "/..")

import openpyxl
from app.boq.writeback import _safe_save

# ---------- 生成测试 BOQ xlsx ----------
wb = openpyxl.Workbook()
ws = wb.active
ws.cell(row=1, column=1, value="BOQ Test")
ws.freeze_panes = "A2"

work_dir = tempfile.mkdtemp(prefix="w5_lock_test_")
boq_path = os.path.join(work_dir, "test_boq.xlsx")
wb.save(boq_path)
print("TEST XLSX:", boq_path)

# ---------- 创建独占锁（Windows API） ----------
GENERIC_READ = 0x80000000
GENERIC_WRITE = 0x40000000
OPEN_EXISTING = 3
FILE_ATTRIBUTE_NORMAL = 0x80
INVALID_HANDLE = ctypes.c_void_p(-1).value

kernel32 = ctypes.windll.kernel32
handle = kernel32.CreateFileW(
    boq_path,
    GENERIC_READ | GENERIC_WRITE,
    0,  # dwShareMode=0：独占
    None,
    OPEN_EXISTING,
    FILE_ATTRIBUTE_NORMAL,
    None,
)
if handle == INVALID_HANDLE:
    print("WARN: CreateFileW failed err=%d — using msvcrt" % ctypes.GetLastError())
    import msvcrt
    fd = os.open(boq_path, os.O_RDWR | os.O_BINARY)
    msvcrt.locking(fd, msvcrt.LK_NBLCK, 1)
    locked_handle = None
    locked_fd = fd
else:
    locked_handle = handle
    locked_fd = None
    print("LOCK: 独占锁已创建 handle=%d" % handle)

# ---------- 直接调 _safe_save（绕过 load_workbook） ----------
# 新建一个 workbook，尝试保存到被锁的路径
wb2 = openpyxl.Workbook()
wb2.active.cell(row=1, column=1, value="test fallback")

try:
    result_path = _safe_save(wb2, boq_path)
    print("SAFE_SAVE result_path:", result_path)

    in_takeoff = "_takeoff" in result_path
    fallback_exists = os.path.exists(result_path)
    print(f"\n{'PASS' if in_takeoff and fallback_exists else 'FAIL'}: "
          f"in_takeoff={in_takeoff} file_exists={fallback_exists}")
except Exception as e:
    print("EXCEPTION:", type(e).__name__, e)
    import traceback; traceback.print_exc()
finally:
    if locked_handle is not None:
        kernel32.CloseHandle(locked_handle)
    elif locked_fd is not None:
        import msvcrt
        msvcrt.locking(locked_fd, msvcrt.LK_UNLCK, 1)
        os.close(locked_fd)
    shutil.rmtree(work_dir, ignore_errors=True)
    print("CLEANUP done")